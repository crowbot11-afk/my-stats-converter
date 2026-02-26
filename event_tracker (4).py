import streamlit as st
import pandas as pd
import pytesseract
from PIL import Image
import io, os, re, json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import date

st.set_page_config(
    page_title="Event Tracker",
    page_icon="swords",
    layout="centered"
)
st.title("Event Participation Tracker")

ROLES = {"deputy", "member", "officer", "elite", "leader", "alliance leader"}


if "ev_loaded" not in st.session_state:
    st.session_state.roster = []
    st.session_state.ev_upload_key = 0
    st.session_state.ev_last_xlsx = None
    st.session_state.ev_matched = {}
    st.session_state.ev_event_name = ""
    st.session_state.ev_loaded = True


def ocr_image(img):
    return pytesseract.image_to_string(img, config="--psm 6")


def alphanum(s):
    return re.sub(r"[^a-z0-9]", "", s.lower())


def clean_name(raw):
    raw = raw.strip()
    raw = re.sub(r"^[^a-zA-Z0-9_\-\.\~\*\+\/]+", "", raw)
    raw = re.sub(r"[^a-zA-Z0-9_\-\.\~\*\+\/\!]+$", "", raw)
    return raw.strip()


def lcs_score(a, b):
    shorter = a if len(a) <= len(b) else b
    longer  = b if len(a) <= len(b) else a
    if len(shorter) < 3:
        return 0
    best = 0
    for start in range(len(shorter)):
        for end in range(start + 3, len(shorter) + 1):
            sub = shorter[start:end]
            if sub in longer and len(sub) > best:
                best = len(sub)
    return best / len(shorter)


def extract_participants(all_text):
    results = {}
    lines = [l.strip() for l in all_text.splitlines()]
    for i, line in enumerate(lines):
        if not line:
            continue
        line_lower = line.lower()
        has_role = any(role in line_lower for role in ROLES)
        if not has_role:
            continue
        not_entered = bool(re.search(r"not\s*entered", line, re.I))
        has_score = bool(re.search(r"\b\d+\b", line))
        if not not_entered and not has_score:
            continue
        player_name = None
        for j in range(i - 1, max(i - 4, -1), -1):
            candidate = lines[j].strip()
            if not candidate:
                continue
            cand_lower = candidate.lower()
            if re.fullmatch(r"[\d,\s\.\-\*\+\=\<\>\(\)\[\]\{\}\/\\]+", candidate):
                continue
            skip = ["battle", "alliance", "ranking", "battlers", "points", "times"]
            if any(w in cand_lower for w in skip):
                continue
            if any(r == cand_lower for r in ROLES):
                continue
            if len(candidate) >= 2 and re.search(r"[a-zA-Z]", candidate):
                player_name = candidate
                break
        if player_name:
            cleaned = clean_name(player_name)
            if len(cleaned) >= 2:
                key = alphanum(cleaned)
                results[key] = (cleaned, not not_entered)
    return results


def fuzzy_match(roster_name, ocr_results, threshold=0.55):
    r_an = alphanum(roster_name)
    if not r_an or len(r_an) < 2:
        return False
    best_score = 0
    best_val = False
    for ocr_key, (ocr_name, entered) in ocr_results.items():
        o_an = alphanum(ocr_key)
        if not o_an:
            continue
        if r_an == o_an:
            return entered
        if len(r_an) >= 4 and len(o_an) >= 4:
            if r_an in o_an or o_an in r_an:
                return entered
        words = [alphanum(w) for w in roster_name.split()]
        words = [w for w in words if len(w) >= 4]
        for word in words:
            if word in o_an or o_an in word:
                score = len(word) / max(len(r_an), 1)
                if score > best_score:
                    best_score = score
                    best_val = entered
        score = lcs_score(r_an, o_an)
        if score > best_score:
            best_score = score
            best_val = entered
    if best_score >= threshold:
        return best_val
    return False


def match_roster(roster, ocr_results):
    matched = {}
    for player in roster:
        matched[player] = fuzzy_match(player, ocr_results)
    return matched


def write_results_to_xlsx(xlsx_file, col_name, results):
    wb = load_workbook(xlsx_file)
    if "Alliance Stats" not in wb.sheetnames:
        return None, "Sheet Alliance Stats not found."
    ws = wb["Alliance Stats"]
    ws.insert_cols(2)
    new_col = 2
    header_cell = ws.cell(1, new_col, col_name)
    header_cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    header_cell.fill = PatternFill("solid", start_color="8B0000")
    header_cell.alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )
    green_fill = PatternFill("solid", start_color="C6EFCE")
    red_fill = PatternFill("solid", start_color="FFC7CE")
    for row in range(2, ws.max_row + 1):
        cell_name = ws.cell(row, 1).value
        if not cell_name:
            continue
        name_str = str(cell_name).strip().lower()
        entered = None
        for rname, val in results.items():
            if rname.strip().lower() == name_str:
                entered = val
                break
        rc = ws.cell(row, new_col)
        rc.alignment = Alignment(horizontal="center", vertical="center")
        if entered is True:
            rc.value = "+"
            rc.fill = green_fill
            rc.font = Font(name="Arial", bold=True, size=11, color="276221")
        else:
            rc.value = "-"
            rc.fill = red_fill
            rc.font = Font(name="Arial", bold=True, size=11, color="9C0006")
    ws.column_dimensions[get_column_letter(new_col)].width = max(len(col_name)+2, 12)
    ws.row_dimensions[1].height = 35
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, None


# =====================
# UI
# =====================
st.subheader("Step 1 - Upload Alliance Tracker")
xlsx_file = st.file_uploader(
    "Upload alliance_tracker.xlsx",
    type=["xlsx"],
    key="ev_xlsx"
)
if xlsx_file and xlsx_file.name != st.session_state.ev_last_xlsx:
    try:
        xl = pd.ExcelFile(xlsx_file)
        df = pd.read_excel(xl, sheet_name="Alliance Stats")
        roster = (
            df["Player Name"].dropna()
            .astype(str).str.strip().tolist()
        )
        roster = [r for r in roster if r]
        st.session_state.roster = roster
        st.session_state.ev_last_xlsx = xlsx_file.name
        st.success("Loaded " + str(len(roster)) + " players.")
    except Exception as e:
        st.error("Error reading file: " + str(e))
if st.session_state.roster:
    with st.expander(
        "View Roster (" + str(len(st.session_state.roster)) + " players)"
    ):
        cols = st.columns(3)
        for i, name in enumerate(st.session_state.roster):
            cols[i % 3].write("- " + name)
st.divider()

st.subheader("Step 2 - Event Name")
today_str = str(date.today())
event_name = st.text_input(
    "Event name (used as column header)",
    value="Elite Wars " + today_str,
    key="ev_name"
)
st.divider()

st.subheader("Step 3 - Upload Screenshots")
screenshots = st.file_uploader(
    "Upload all event leaderboard screenshots",
    type=["png", "jpg", "jpeg", "webp"],
    accept_multiple_files=True,
    key="ev_shots_" + str(st.session_state.ev_upload_key)
)
has_roster = bool(st.session_state.roster)
has_shots = bool(screenshots)
has_name = bool(event_name.strip())
can_process = has_roster and has_shots and has_name

col_a, col_b = st.columns([2, 1])
with col_a:
    do_process = st.button(
        "Process Event",
        disabled=not can_process,
        type="primary",
        use_container_width=True
    )
with col_b:
    if st.button("Clear", use_container_width=True):
        st.session_state.ev_upload_key += 1
        st.session_state.ev_matched = {}
        st.rerun()
if not has_roster:
    st.caption("Upload the alliance tracker file first.")
elif not has_name:
    st.caption("Enter an event name.")
elif not has_shots:
    st.caption("Upload at least one screenshot.")

if do_process and can_process:
    all_text = ""
    bar = st.progress(0)
    slot = st.empty()
    for i, f in enumerate(screenshots):
        slot.write(
            "Reading " + str(i+1) + "/" + str(len(screenshots)) + "..."
        )
        all_text += chr(10) + ocr_image(Image.open(f))
        bar.progress((i+1) / len(screenshots))
    slot.empty()
    bar.empty()
    ocr_results = extract_participants(all_text)
    matched = match_roster(st.session_state.roster, ocr_results)
    st.session_state.ev_matched = matched
    st.session_state.ev_event_name = event_name.strip()
    entered_n = sum(1 for v in matched.values() if v)
    not_entered_n = sum(1 for v in matched.values() if not v)
    st.success(
        "Done! " + str(entered_n) + " entered, "
        + str(not_entered_n) + " not entered."
    )
    with st.expander("OCR detected names (debug)"):
        for k, (n, e) in ocr_results.items():
            status = "ENTERED" if e else "NOT ENTERED"
            st.write("- " + n + " : " + status)
    with st.expander("Raw OCR text (debug)"):
        st.text_area("", all_text, height=200)

if st.session_state.ev_matched:
    st.divider()
    st.subheader("Results")
    matched = st.session_state.ev_matched
    entered_rows = [{"Player": n} for n, v in matched.items() if v]
    not_entered_rows = [{"Player": n} for n, v in matched.items() if not v]
    col1, col2 = st.columns(2)
    with col1:
        st.markdown("**Entered (" + str(len(entered_rows)) + ")**")
        if entered_rows:
            st.dataframe(
                pd.DataFrame(entered_rows),
                use_container_width=True,
                hide_index=True
            )
    with col2:
        st.markdown("**Not Entered (" + str(len(not_entered_rows)) + ")**")
        if not_entered_rows:
            st.dataframe(
                pd.DataFrame(not_entered_rows),
                use_container_width=True,
                hide_index=True
            )
    st.divider()
    st.subheader("Step 4 - Download Updated Tracker")
    if xlsx_file is not None:
        xlsx_file.seek(0)
        buf, err = write_results_to_xlsx(
            xlsx_file,
            st.session_state.ev_event_name,
            matched
        )
        if err:
            st.error(err)
        else:
            st.download_button(
                label="Download Updated alliance_tracker.xlsx",
                data=buf,
                file_name="alliance_tracker.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
    else:
        st.warning("Re-upload the tracker above to save results.")
