import streamlit as st
import pandas as pd
import pytesseract
from PIL import Image
import io, os, json
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

st.set_page_config(page_title="Alliance Tracker", page_icon="sword", layout="centered")
st.title("Alliance Tracker")

KNOWN_STATS = [
    "Infantry Attack",
    "Infantry Defense",
    "Infantry HP",
    "Cavalry Attack",
    "Cavalry Defense",
    "Cavalry HP",
    "Archer Attack",
    "Archer Defense",
    "Archer HP",
    "Mage Attack",
    "Mage Defense",
    "Mage HP",
    "Angel Attack",
    "Angel Defense",
    "Angel HP",
    "Golem Attack",
    "Golem Defense",
    "Golem HP",
    "Enemy Troops Attack Reduction",
    "Enemy Troops HP Reduction",
    "Archer Damage",
    "Mage Damage",
    "Troops Damage Taken Reduction",
    "Damage Boost when attacking",
    "Damage taken reduced when attacking",
    "Damage Boost when defending",
    "Damage taken reduced when defending",
    "Infantry Resilience",
    "Cavalry Resilience",
    "Archer Penetration",
    "Mage Penetration",
    "Archer Mastery",
    "Mage Mastery",
    "Damage Against Infantry Boost",
    "Damage Against Cavalry Boost",
    "Damage Against Angels Boost",
    "Infantry Damage Taken Reduction",
    "Cavalry Damage Taken Reduction",
    "Reduces Damage taken from Infantry",
    "Reduces Damage taken from Cavalry",
    "Reduces Damage taken from Archers",
    "Reduces Damage taken from Mages",
    "Mage damage increased on Infantry",
    "Mage damage increased on Cavalry",
    "Mage damage increased on Archers",
    "Critical Strike Rate Boost",
    "Critical Strike Rate Taken Reduction",
]

ALL_COLUMNS = ["Player Name", "March Size"] + KNOWN_STATS + [
    "Elder Titan Tier", "Titan Talent Level", "Beast Tier", "Beast Talent Level",
    "Beast Skill Level", "Totem Level", "Special Stats Level", "Jewels Level",
    "Zodiac Green", "Zodiac White", "Northern Green", "Colossus Level", "Emblem Level",
    "Frontline", "Backline"
]

TOTAL_ROWS = 100
DATA_FILE = "/tmp/alliance_data.json"


def persist():
    payload = {"roster": st.session_state.roster, "df": st.session_state.df.to_dict("records")}
    with open(DATA_FILE, "w") as f:
        json.dump(payload, f)


def load_persisted():
    if os.path.exists(DATA_FILE):
        try:
            with open(DATA_FILE, "r") as f:
                payload = json.load(f)
            roster = payload.get("roster", [])
            records = payload.get("df", [])
            df = pd.DataFrame(records, columns=ALL_COLUMNS) if records else pd.DataFrame(columns=ALL_COLUMNS)
            for col in ALL_COLUMNS:
                if col not in df.columns:
                    df[col] = 0
            return roster, df
        except Exception:
            pass
    return [], pd.DataFrame(columns=ALL_COLUMNS)


if "loaded" not in st.session_state:
    roster, df = load_persisted()
    st.session_state.roster = roster
    st.session_state.df = df
    st.session_state.extracted = None
    st.session_state.upload_key = 0
    st.session_state.add_msg = None
    st.session_state.kick_msg = None
    st.session_state.save_msg = None
    st.session_state.chosen_player = ""
    st.session_state.last_loaded_file = None
    st.session_state.frontline = "Infantry"
    st.session_state.backline = "Archer"
    st.session_state.loaded = True


def ocr_image(img):
    return pytesseract.image_to_string(img, config="--psm 6")


def extract_all(all_text):
    data = {}
    m = re.search(r"Total Army\s+([\d,]+)", all_text)
    if m:
        data["March Size"] = m.group(1).replace(",", "")
    for stat in KNOWN_STATS:
        pattern = re.escape(stat).replace(r"\ ", r"\s+")
        m = re.search(pattern + r"\s+([\d,]+\.?\d*%?)", all_text, re.I)
        if m:
            data[stat] = m.group(1).strip()
    m = re.search(r"Evolution:\s*Titan Tier\s*(III|II|I|lll|ll|l|\d)", all_text, re.I)
    if m:
        tier = m.group(1).replace("lll", "III").replace("ll", "II").replace("l", "I")
        data["Elder Titan Tier"] = "Titan Tier " + tier
    talent_levels = re.findall(r"Total Talent Level:\s*(\d+)", all_text)
    if talent_levels:
        data["Titan Talent Level"] = talent_levels[0]
    m = re.search(r"[Ee]volution:\s*Tier\s*(\d+)", all_text)
    if m:
        data["Beast Tier"] = m.group(1)
    if len(talent_levels) >= 2:
        data["Beast Talent Level"] = talent_levels[1]
    m = re.search(r"Total [Ss]kill Level:\s*(\d+)", all_text)
    if m:
        data["Beast Skill Level"] = m.group(1)
    m = re.search(r"(?:Level|Lv)[:\s.]*\s*(Lv\.?\s*\d+)", all_text, re.I)
    if m:
        data["Totem Level"] = re.sub(r"\s+", "", m.group(1))
    m = re.search(r"Total Special Stats Level:\s*(\d+)", all_text)
    if m:
        data["Special Stats Level"] = m.group(1)
    m = re.search(r"Total Jewels Level:\s*(\d+)", all_text)
    if m:
        data["Jewels Level"] = m.group(1)
    m = re.search(r"Zodiac Palace.*?Green Amount:\s*(\d+).*?White Amount:\s*(\d+)", all_text, re.S)
    if m:
        data["Zodiac Green"] = m.group(1)
        data["Zodiac White"] = m.group(2)
    m = re.search(r"Northern Palace.*?Green Amount:\s*(\d+)", all_text, re.S)
    if m:
        data["Northern Green"] = m.group(1)
    m = re.search(r"Colossus.*?Total Level:\s*(\d+)", all_text, re.S)
    if m:
        data["Colossus Level"] = m.group(1)
    m = re.search(r"Emblem.*?Total Level:\s*(\d+)", all_text, re.S)
    if m:
        data["Emblem Level"] = m.group(1)
    return data


def save_player(name, data):
    data["Player Name"] = name
    df = st.session_state.df.copy()
    name_lower = name.strip().lower()
    mask = df["Player Name"].astype(str).str.strip().str.lower() == name_lower
    action = "updated" if mask.any() else "added"
    df = df[~mask].reset_index(drop=True)
    new_row = {col: data.get(col, 0) for col in ALL_COLUMNS}
    new_df = pd.DataFrame([new_row], columns=ALL_COLUMNS)
    df = pd.concat([df, new_df], ignore_index=True)
    st.session_state.df = df
    return action


def build_excel():
    df = st.session_state.df.copy()
    roster = sorted(st.session_state.roster)
    for col in ALL_COLUMNS:
        if col not in df.columns:
            df[col] = 0
    df = df.reindex(columns=ALL_COLUMNS)
    wb = Workbook()
    ws_r = wb.active
    ws_r.title = "Roster"
    ws_r["A1"] = "Player Names"
    for i, name in enumerate(roster, start=2):
        ws_r.cell(i, 1).value = name
    ws_r.sheet_state = "hidden"
    ws = wb.create_sheet("Alliance Stats")
    h_fill = PatternFill("solid", start_color="8B0000")
    h_font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    for ci, col in enumerate(ALL_COLUMNS, 1):
        c = ws.cell(1, ci, col)
        c.fill = h_fill
        c.font = h_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    alt = PatternFill("solid", start_color="FFF3E0")
    empty_fill = PatternFill("solid", start_color="F9F9F9")
    records = df.to_dict("records")
    for ri in range(TOTAL_ROWS):
        er = ri + 2
        rec = records[ri] if ri < len(records) else {}
        has_data = bool(rec)
        for ci, col in enumerate(ALL_COLUMNS, 1):
            c = ws.cell(er, ci)
            c.value = rec.get(col) if has_data else None
            c.font = Font(name="Arial", size=10)
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.fill = (alt if ri % 2 == 0 else PatternFill()) if has_data else empty_fill
    if roster:
        dv = DataValidation(
            type="list",
            formula1="Roster!$A$2:$A$" + str(len(roster) + 1),
            allow_blank=True,
            showDropDown=False
        )
        dv.sqref = "A2:A" + str(TOTAL_ROWS + 1)
        ws.add_data_validation(dv)
    ws.column_dimensions["A"].width = 22
    for ci, col in enumerate(ALL_COLUMNS, 1):
        if ci > 1:
            ws.column_dimensions[get_column_letter(ci)].width = max(len(col) + 2, 12)
    ws.row_dimensions[1].height = 35
    ws.freeze_panes = "B2"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def load_excel(file):
    xl = pd.ExcelFile(file)
    df = pd.read_excel(xl, sheet_name="Alliance Stats")
    df = df[df["Player Name"].notna() & (df["Player Name"].astype(str).str.strip() != "")]
    df = df.reset_index(drop=True)
    try:
        rdf = pd.read_excel(xl, sheet_name="Roster", header=0)
        roster = rdf.iloc[:, 0].dropna().astype(str).str.strip().tolist()
    except Exception:
        roster = df["Player Name"].dropna().astype(str).str.strip().tolist()
    return df, roster


tab1, tab2, tab3 = st.tabs(["Load & Roster", "Add / Update Player", "Download"])

# TAB 1
with tab1:
    if st.session_state.add_msg:
        lvl, txt = st.session_state.add_msg
        st.session_state.add_msg = None
        st.success(txt) if lvl == "success" else st.warning(txt)
    if st.session_state.kick_msg:
        lvl, txt = st.session_state.kick_msg
        st.session_state.kick_msg = None
        st.success(txt) if lvl == "success" else st.warning(txt)
    st.subheader("Load Existing Tracker")
    uploaded = st.file_uploader("Upload alliance_tracker.xlsx", type=["xlsx"], key="load")
    if uploaded and uploaded.name != st.session_state.last_loaded_file:
        df, roster = load_excel(uploaded)
        st.session_state.df = df
        st.session_state.roster = roster
        st.session_state.chosen_player = ""
        st.session_state.last_loaded_file = uploaded.name
        persist()
        st.success("Loaded " + str(len(df)) + " players, " + str(len(roster)) + " in roster.")
    st.info(str(len(st.session_state.df)) + " players with data / " + str(len(st.session_state.roster)) + " in roster")
    st.divider()
    st.subheader("Roster Management")
    col_add, col_kick = st.columns(2)
    with col_add:
        st.markdown("**Add Player**")
        st.text_input("Player name", placeholder="Type name here", key="inp_new_name")
        if st.button("Add to Roster", use_container_width=True):
            name = st.session_state.inp_new_name.strip()
            if not name:
                st.session_state.add_msg = ("warning", "Type a name first.")
            elif any(name.lower() == n.lower() for n in st.session_state.roster):
                st.session_state.add_msg = ("warning", name + " is already in the roster.")
            else:
                st.session_state.roster = sorted(st.session_state.roster + [name])
                persist()
                st.session_state.add_msg = ("success", name + " added to roster!")
            st.rerun()
    with col_kick:
        st.markdown("**Remove Player**")
        if st.session_state.roster:
            to_kick = st.selectbox("Player to remove", ["-- select --"] + st.session_state.roster, label_visibility="collapsed")
            if st.button("Remove and Delete Data", use_container_width=True):
                if to_kick == "-- select --":
                    st.session_state.kick_msg = ("warning", "Select a player first.")
                else:
                    st.session_state.roster = [n for n in st.session_state.roster if n.lower() != to_kick.lower()]
                    st.session_state.df = st.session_state.df[
                        st.session_state.df["Player Name"].astype(str).str.strip().str.lower() != to_kick.lower()
                    ].reset_index(drop=True)
                    if st.session_state.chosen_player.lower() == to_kick.lower():
                        st.session_state.chosen_player = ""
                    persist()
                    st.session_state.kick_msg = ("success", to_kick + " removed.")
                st.rerun()
        else:
            st.info("No players in roster yet.")
    if st.session_state.roster:
        st.divider()
        with st.expander("View Full Roster (" + str(len(st.session_state.roster)) + " players)"):
            cols = st.columns(3)
            for i, name in enumerate(st.session_state.roster):
                cols[i % 3].write("- " + name)

# TAB 2
with tab2:
    if st.session_state.save_msg:
        lvl, txt = st.session_state.save_msg
        st.session_state.save_msg = None
        st.success(txt) if lvl == "success" else st.warning(txt)
    if not st.session_state.roster:
        st.info("Go to Load & Roster tab and add players first.")
    else:
        st.subheader("Select Player")
        _opts = ["-- select player --"] + sorted(st.session_state.roster)
        if st.session_state.chosen_player in st.session_state.roster:
            _cur_idx = _opts.index(st.session_state.chosen_player)
        else:
            _cur_idx = 0
        selected = st.selectbox("Select player", options=_opts, index=_cur_idx, key="_sel_player_box")
        if selected == "-- select player --":
            st.session_state.chosen_player = ""
        else:
            st.session_state.chosen_player = selected
        player_name = st.session_state.chosen_player
        if player_name:
            df_check = st.session_state.df
            mask = df_check["Player Name"].astype(str).str.strip().str.lower() == player_name.strip().lower()
            if mask.any():
                row = df_check[mask].iloc[0]
                stat_cols = [c for c in ALL_COLUMNS if c != "Player Name"]
                has_real_data = any(
                    str(row.get(c, 0)) not in ("0", "0.0", "", "nan", "None")
                    for c in stat_cols
                )
                if has_real_data:
                    st.warning(player_name + " already has data - saving will replace it.")
                else:
                    st.info("Selected: " + player_name + " - ready for first upload.")
            else:
                st.info("Selected: " + player_name)
        st.divider()
        st.subheader("Role Selection")
        role_col1, role_col2 = st.columns(2)
        with role_col1:
            st.markdown("**Frontline**")
            frontline = st.radio(
                "Frontline",
                options=["Infantry", "Cavalry"],
                index=0 if st.session_state.frontline == "Infantry" else 1,
                horizontal=True,
                key="frontline_radio",
                label_visibility="collapsed"
            )
            st.session_state.frontline = frontline
        with role_col2:
            st.markdown("**Backline**")
            backline = st.radio(
                "Backline",
                options=["Archer", "Mage"],
                index=0 if st.session_state.backline == "Archer" else 1,
                horizontal=True,
                key="backline_radio",
                label_visibility="collapsed"
            )
            st.session_state.backline = backline
        st.divider()
        st.subheader("Upload Screenshots")
        screenshots = st.file_uploader(
            "Select all screenshots at once",
            type=["png", "jpg", "jpeg", "webp"],
            accept_multiple_files=True,
            key="shots_" + str(st.session_state.upload_key)
        )
        has_player = bool(player_name)
        has_shots = bool(screenshots)
        can_extract = has_player and has_shots
        col_a, col_b = st.columns([2, 1])
        with col_a:
            do_extract = st.button("Extract Stats", disabled=not can_extract, use_container_width=True)
        with col_b:
            if st.button("Clear", use_container_width=True):
                st.session_state.upload_key += 1
                st.session_state.extracted = None
                st.rerun()
        if not has_player:
            st.caption("Select a player above to enable extraction.")
        elif not has_shots:
            st.caption("Upload at least one screenshot to enable extraction.")
        if do_extract and can_extract:
            all_text = ""
            bar = st.progress(0)
            slot = st.empty()
            for i, f in enumerate(screenshots):
                slot.write("Reading " + str(i + 1) + "/" + str(len(screenshots)) + "...")
                all_text += chr(10) + ocr_image(Image.open(f))
                bar.progress((i + 1) / len(screenshots))
            slot.empty()
            bar.empty()
            st.session_state.extracted = extract_all(all_text)
            st.session_state.extracted["_player"] = player_name
            n = len([k for k in st.session_state.extracted if k != "_player"])
            st.success(str(n) + " fields extracted for " + player_name)
            stats_found = {k: st.session_state.extracted[k] for k in KNOWN_STATS if k in st.session_state.extracted}
            build_keys = [
                "March Size", "Elder Titan Tier", "Titan Talent Level", "Beast Tier",
                "Beast Talent Level", "Beast Skill Level", "Totem Level",
                "Special Stats Level", "Jewels Level", "Zodiac Green", "Zodiac White",
                "Northern Green", "Colossus Level", "Emblem Level"
            ]
            build_found = {k: st.session_state.extracted[k] for k in build_keys if k in st.session_state.extracted}
            st.markdown("**Build Info**")
            if build_found:
                st.dataframe(pd.DataFrame(list(build_found.items()), columns=["Field", "Value"]), use_container_width=True)
            else:
                st.caption("No build info found in screenshots.")
            st.markdown("**Stats Bonus**")
            if stats_found:
                items = list(stats_found.items())
                half = len(items) // 2 + len(items) % 2
                ca, cb = st.columns(2)
                with ca:
                    st.dataframe(pd.DataFrame(items[:half], columns=["Stat", "Value"]), use_container_width=True)
                with cb:
                    st.dataframe(pd.DataFrame(items[half:], columns=["Stat", "Value"]), use_container_width=True)
                st.info(str(len(stats_found)) + "/47 stats found")
            else:
                st.caption("No combat stats found in screenshots.")
            with st.expander("Raw OCR"):
                st.text_area("", all_text, height=150)
        if (
            st.session_state.extracted is not None
            and st.session_state.extracted.get("_player") == player_name
            and player_name
        ):
            st.divider()
            df_check = st.session_state.df
            mask = df_check["Player Name"].astype(str).str.strip().str.lower() == player_name.strip().lower()
            has_real_data = False
            if mask.any():
                row = df_check[mask].iloc[0]
                stat_cols = [c for c in ALL_COLUMNS if c != "Player Name"]
                has_real_data = any(
                    str(row.get(c, 0)) not in ("0", "0.0", "", "nan", "None")
                    for c in stat_cols
                )
            label = "Update " + player_name if has_real_data else "Save " + player_name
            if st.button(label, type="primary", use_container_width=True):
                data_to_save = {k: v for k, v in st.session_state.extracted.items() if k != "_player"}
                data_to_save["Frontline"] = st.session_state.frontline
                data_to_save["Backline"] = st.session_state.backline
                action = save_player(player_name, data_to_save)
                persist()
                df_verify = st.session_state.df
                v_mask = df_verify["Player Name"].astype(str).str.strip().str.lower() == player_name.strip().lower()
                if v_mask.any():
                    saved_fields = sum(
                        1 for c in ALL_COLUMNS if c != "Player Name"
                        and str(df_verify[v_mask].iloc[0].get(c, 0)) not in ("0", "0.0", "", "nan", "None")
                    )
                    verb = "Updated" if action == "updated" else "Saved"
                    st.session_state.save_msg = ("success", verb + " " + player_name + " - " + str(saved_fields) + " fields stored.")
                else:
                    st.session_state.save_msg = ("warning", "Save failed for " + player_name + ". Please try again.")
                st.session_state.extracted = None
                st.session_state.upload_key += 1
                st.rerun()

# TAB 3
with tab3:
    st.subheader("Download Alliance Tracker")
    st.write(
        str(len(st.session_state.df)) + " players with data / "
        + str(len(st.session_state.roster)) + " in roster / "
        + str(TOTAL_ROWS) + " rows in Excel"
    )
    if len(st.session_state.df) > 0:
        with st.expander("Preview data"):
            st.dataframe(st.session_state.df, use_container_width=True)
    excel_buf = build_excel()
    st.download_button(
        label="Download alliance_tracker.xlsx",
        data=excel_buf,
        file_name="alliance_tracker.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        use_container_width=True
    )
    st.caption("Save to Google Drive. Next session upload it here to continue.")
