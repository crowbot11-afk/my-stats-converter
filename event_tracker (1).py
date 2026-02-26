import streamlit as st
import pandas as pd
import pytesseract
from PIL import Image
import io, os, re, json
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from datetime import date

st.set_page_config(
    page_title="Event Tracker",
    page_icon="swords",
    layout="centered"
)
st.title("Event Participation Tracker")

DATA_FILE = "/tmp/event_tracker.json"
ROLES = {"deputy", "member", "officer", "elite", "leader"}


if "ev_loaded" not in st.session_state:
    st.session_state.roster = []
    st.session_state.ev_upload_key = 0
    st.session_state.ev_last_xlsx = None
    st.session_state.ev_matched = {}
    st.session_state.ev_event_name = ""
    st.session_state.ev_loaded = True


def ocr_image(img):
    return pytesseract.image_to_string(img, config="--psm 6")


def clean_name(raw):
    raw = raw.strip()
    raw = re.sub(r"^[^a-zA-Z0-9_\-\.\~\*\+\/]+", "", raw)
    raw = re.sub(r"[^a-zA-Z0-9_\-\.\~\*\+\/\!]+$", "", raw)
    return raw.strip()


def extract_participants(all_text):
    results = {}
    lines = [l.strip() for l in all_text.splitlines()]
    for i, line in enumerate(lines):
        if not line:
            continue
        line_lower = line.lower()
        has_role = any(role in line_lower for role in ROLES)
        if not has_role:
            continue
        not_entered = bool(re.search(r"not\s*entered", line, re.I))
        has_score = bool(re.search(r"\b\d+\b", line))
        if not not_entered and not has_score:
            continue
        player_name = None
        for j in range(i - 1, max(i - 4, -1), -1):
            candidate = lines[j].strip()
            if not candidate:
                continue
            cand_lower = candidate.lower()
            if re.fullmatch(r"[\d,\s\.\-\*\+\=\<\>\(\)\[\]\{\}\/\\]+", candidate):
                continue
            skip = ["battle", "alliance", "ranking", "battlers", "points", "times"]
            if any(w in cand_lower for w in skip):
                continue
            if any(r == cand_lower for r in ROLES):
                continue
            if len(candidate) >= 2 and re.search(r"[a-zA-Z]", candidate):
                player_name = candidate
                break
        if player_name:
            cleaned = clean_name(player_name)
            if len(cleaned) >= 2:
                key = cleaned.lower()
                results[key] = (cleaned, not not_entered)
    return results


def fuzzy_match(roster_name, ocr_results):
    r_lower = roster_name.strip().lower()
    r_clean = re.sub(r"[^a-z0-9]", "", r_lower)
    if r_lower in ocr_results:
        return ocr_results[r_lower][1]
    for ocr_key, (ocr_name, entered) in ocr_results.items():
        ocr_clean = re.sub(r"[^a-z0-9]", "", ocr_key)
        if r_clean == ocr_clean and len(r_clean) >= 3:
            return entered
    for ocr_key, (ocr_name, entered) in ocr_results.items():
        ocr_clean = re.sub(r"[^a-z0-9]", "", ocr_key)
        if len(r_clean) >= 4 and (r_clean in ocr_clean or ocr_clean in r_clean):
            return entered
    return False


def match_roster(roster, ocr_results):
    matched = {}
    for player in roster:
        matched[player] = fuzzy_match(player, ocr_results)
    return matched


def write_results_to_xlsx(xlsx_file, col_name, results):
    wb = load_workbook(xlsx_file)
    if "Alliance Stats" not in wb.sheetnames:
        return None, "Sheet Alliance Stats not found."
    ws = wb["Alliance Stats"]
    new_col = ws.max_column + 1
    header_cell = ws.cell(1, new_col, col_name)
    header_cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
    header_cell.fill = PatternFill("solid", start_color="8B0000")
    header_cell.alignment = Alignment(
        horizontal="center", vertical="center", wrap_text=True
    )
    green_fill = PatternFill("solid", start_color="C6EFCE")
    red_fill = PatternFill("solid", start_color="FFC7CE")
    for row in range(2, ws.max_row + 1):
        cell_name = ws.cell(row, 1).value
        if not cell_name:
            continue
        name_str = str(cell_name).strip().lower()
        entered = None
        for rname, val in results.items():
            if rname.strip().lower() == name_str:
                entered = val
                break
        rc = ws.cell(row, new_col)
        rc.alignment = Alignment(horizontal="center", vertical="center")
        if entered is True:
            rc.value = "+"
            rc.fill = green_fill
            rc.font = Font(name="Arial", bold=True, size=11, color="276221")
        elif entered is False:
            rc.value = "-"
            rc.fill = red_fill
            rc.font = Font(name="Arial", bold=True, size=11, color="9C0006")
    ws.column_dimensions[get_column_letter(new_col)].width = max(len(col_name)+2, 12)
    ws.row_dimensions[1].height = 35
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf, None


# =====================
# UI
# =====================
st.subheader("Step 1 - Upload Alliance Tracker")
xlsx_file = st.file_uploader(
    "Upload alliance_tracker.xlsx",
    type=["xlsx"],
    key="ev_xlsx"
)
if xlsx_file and xlsx_file.name != st.session_state.ev_last_xlsx:
    try:
        xl = pd.ExcelFile(xlsx_file)
        df = pd.read_excel(xl, sheet_name="Alliance Stats")
        roster = (
            df["Player Name"].dropna()
            .astype(str).str.strip().tolist()
        )
        roster = [r for r in roster if r]
        st.session_state.roster = roster
        st.session_state.ev_last_xlsx = xlsx_file.name
        st.success("Loaded " + str(len(roster)) + " players.")
    except Exception as e:
        st.error("Error reading file: " + str(e))
if st.session_state.roster:
    with st.expander(
        "View Roster (" + str(len(st.session_state.roster)) + " players)"
    ):
        cols = st.columns(3)
        for i, name in enumerate(st.session_state.roster):
            cols[i % 3].write("- " + name)
st.divider()

st.subheader("Step 2 - Event Name")
today_str = str(date.today())
event_name = st.text_input(
    "Event name (used as column header)",
    value="Elite Wars " + today_str,
    key="ev_name"
)
st.divider()

st.subheader("Step 3 - Upload Screenshots")
screenshots = st.file_uploader(
    "Upload all event leaderboard screenshots",
    type=["png", "jpg", "jpeg", "webp"],
    accept_multiple_files=True,
    key="ev_shots_" + str(st.session_state.ev_upload_key)
)
has_roster = bool(st.session_state.roster)
has_shots = bool(screenshots)
has_name = bool(event_name.strip())
can_process = has_roster and has_shots and has_name

col_a, col_b = st.columns([2, 1])
with col_a:
    do_process = st.button(
        "Process Event",
        disabled=not can_process,
        type="primary",
        use_container_width=True
    )
with col_b:
    if st.button("Clear", use_container_width=True):
        st.session_state.ev_upload_key += 1
        st.session_state.ev_matched = {}
        st.rerun()
if not has_roster:
    st.caption("Upload the alliance tracker file first.")
elif not has_name:
    st.caption("Enter an event name.")
elif not has_shots:
    st.caption("Upload at least one screenshot.")

if do_process and can_process:
    all_text = ""
    bar = st.progress(0)
    slot = st.empty()
    for i, f in enumerate(screenshots):
        slot.write(
            "Reading " + str(i+1) + "/" + str(len(screenshots)) + "..."
        )
        all_text += chr(10) + ocr_image(Image.open(f))
        bar.progress((i+1) / len(screenshots))
    slot.empty()
    bar.empty()
    ocr_results = extract_participants(all_text)
    matched = match_roster(st.session_state.roster, ocr_results)
    st.session_state.ev_matched = matched
    st.session_state.ev_event_name = event_name.strip()
    entered_n = sum(1 for v in matched.values() if v)
    not_entered_n = sum(1 for v in matched.values() if not v)
    st.success(
        "Done! " + str(entered_n) + " entered, "
        + str(not_entered_n) + " not entered."
    )
    with st.expander("Raw OCR (debug)"):
        st.text_area("", all_text, height=200)
    with st.expander("OCR detected names (debug)"):
        for k, (n, e) in ocr_results.items():
            status = "ENTERED" if e else "NOT ENTERED"
            st.write("- " + n + " : " + status)

if st.session_state.ev_matched:
    st.divider()
    st.subheader("Results")
    matched = st.session_state.ev_matched
    entered_rows = [{"Player": n} for n, v in matched.items() if v]
    not_entered_rows = [{"Player": n} for n, v in matched.items() if not v]
    col1, col2 = st.columns(2)
    with col1:
        st.markdown("**Entered (" + str(len(entered_rows)) + ")**")
        if entered_rows:
            st.dataframe(
                pd.DataFrame(entered_rows),
                use_container_width=True,
                hide_index=True
            )
    with col2:
        st.markdown("**Not Entered (" + str(len(not_entered_rows)) + ")**")
        if not_entered_rows:
            st.dataframe(
                pd.DataFrame(not_entered_rows),
                use_container_width=True,
                hide_index=True
            )
    st.divider()
    st.subheader("Step 4 - Download Updated Tracker")
    if xlsx_file is not None:
        xlsx_file.seek(0)
        buf, err = write_results_to_xlsx(
            xlsx_file,
            st.session_state.ev_event_name,
            matched
        )
        if err:
            st.error(err)
        else:
            st.download_button(
                label="Download Updated alliance_tracker.xlsx",
                data=buf,
                file_name="alliance_tracker.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True
            )
    else:
        st.warning("Re-upload the tracker above to save results.")
