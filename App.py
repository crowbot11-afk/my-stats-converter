import streamlit as st
import pandas as pd
import pytesseract
from PIL import Image
import io
import re
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

st.set_page_config(page_title="Alliance Battle Tracker", page_icon="⚔️", layout="centered")
st.title("⚔️ Alliance Battle Report Tracker")
st.write("Upload all screenshots for one player → extract stats → add to Alliance Excel.")

KNOWN_STATS = [
    "Infantry Attack", "Infantry Defense", "Infantry HP",
    "Cavalry Attack", "Cavalry Defense", "Cavalry HP",
    "Archer Attack", "Archer Defense", "Archer HP",
    "Mage Attack", "Mage Defense", "Mage HP",
    "Angel Attack", "Angel Defense", "Angel HP",
    "Golem Attack", "Golem Defense", "Golem HP",
    "Enemy Troops Attack Reduction", "Enemy Troops HP Reduction",
    "Archer Damage", "Mage Damage", "Troops Damage Taken Reduction",
    "Damage Boost when attacking", "Damage taken reduced when attacking",
    "Damage Boost when defending", "Damage taken reduced when defending",
    "Infantry Resilience", "Cavalry Resilience",
    "Archer Penetration", "Mage Penetration",
    "Archer Mastery", "Mage Mastery",
    "Damage Against Infantry Boost", "Damage Against Cavalry Boost",
    "Damage Against Angels Boost",
    "Infantry Damage Taken Reduction", "Cavalry Damage Taken Reduction",
    "Reduces Damage taken from Infantry", "Reduces Damage taken from Cavalry",
    "Reduces Damage taken from Archers", "Reduces Damage taken from Mages",
    "Mage damage increased on Infantry", "Mage damage increased on Cavalry",
    "Mage damage increased on Archers",
    "Critical Strike Rate Boost", "Critical Strike Rate Taken Reduction",
]

def ocr_image(img):
    return pytesseract.image_to_string(img, config='--psm 6')

def extract_all(all_text):
    data = {}

    # Player name: OCR renders ^^Name^^ as ""Name"" or ""Name™ etc
    m = re.search(r'[""]{1,2}([A-Za-z0-9_~\-. ^]+?)[""™]{1,2}', all_text)
    if m:
        name = m.group(1).strip().strip('^').strip()
        if len(name) > 1:
            data['Player Name'] = name

    # Result
    m = re.search(r'\b(Victory|Defeat)\b', all_text, re.I)
    if m:
        data['Result'] = m.group(1)

    # Power Loss (take the attacker's, which is first)
    m = re.search(r'Power\s*Loss\s+([-\d,]+)', all_text)
    if m:
        data['Power Loss'] = m.group(1).replace(',', '')

    # Battle stats (first occurrence = attacker)
    for stat in ['Total Army', 'Kills', 'Losses', 'Wounded', 'Survivors', 'Battlers']:
        m = re.search(rf'{stat}\s+([\d,]+)', all_text)
        if m:
            data[stat] = m.group(1).replace(',', '')

    # All Stats Bonus - match each known stat precisely
    for stat in KNOWN_STATS:
        pattern = re.escape(stat).replace(r'\ ', r'\s+')
        m = re.search(pattern + r'\s+([\d,]+\.?\d*%?)', all_text, re.I)
        if m:
            data[stat] = m.group(1).strip()

    # Elder Titan tier
    m = re.search(r'Evolution:\s*Titan Tier\s*(III|II|I|lll|ll|l|\d)', all_text, re.I)
    if m:
        tier = m.group(1).replace('lll','III').replace('ll','II').replace('l','I')
        data['Elder Titan Tier'] = 'Titan Tier ' + tier
    talent_levels = re.findall(r'Total Talent Level:\s*(\d+)', all_text)
    if talent_levels:
        data['Titan Talent Level'] = talent_levels[0]

    # Beast
    m = re.search(r'[Ee]volution:\s*Tier\s*(\d+)', all_text)
    if m:
        data['Beast Tier'] = m.group(1)
    if len(talent_levels) >= 2:
        data['Beast Talent Level'] = talent_levels[1]
    m = re.search(r'Total [Ss]kill Level:\s*(\d+)', all_text)
    if m:
        data['Beast Skill Level'] = m.group(1)

    # Totem
    m = re.search(r'(?:Level|Lv)[:\s.]*\s*(Lv\.?\s*\d+)', all_text, re.I)
    if m:
        data['Totem Level'] = re.sub(r'\s+', '', m.group(1))

    # Equip
    m = re.search(r'Total Special Stats Level:\s*(\d+)', all_text)
    if m:
        data['Special Stats Level'] = m.group(1)
    m = re.search(r'Total Jewels Level:\s*(\d+)', all_text)
    if m:
        data['Jewels Level'] = m.group(1)

    # Star Palace
    m = re.search(r'Zodiac Palace.*?Green Amount:\s*(\d+).*?White Amount:\s*(\d+)', all_text, re.S)
    if m:
        data['Zodiac Green'] = m.group(1)
        data['Zodiac White'] = m.group(2)
    m = re.search(r'Northern Palace.*?Green Amount:\s*(\d+)', all_text, re.S)
    if m:
        data['Northern Green'] = m.group(1)

    # Colossus
    m = re.search(r'Colossus.*?Total Level:\s*(\d+)', all_text, re.S)
    if m:
        data['Colossus Level'] = m.group(1)

    # Emblem
    m = re.search(r'Emblem.*?Total Level:\s*(\d+)', all_text, re.S)
    if m:
        data['Emblem Level'] = m.group(1)

    return data

def build_excel(df):
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name='Alliance Stats')
        ws = writer.sheets['Alliance Stats']

        header_fill = PatternFill('solid', start_color='8B0000')
        header_font = Font(name='Arial', bold=True, color='FFFFFF', size=10)
        alt_fill = PatternFill('solid', start_color='FFF3E0')

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

        for i, row in enumerate(ws.iter_rows(min_row=2), start=2):
            for cell in row:
                cell.font = Font(name='Arial', size=10)
                cell.alignment = Alignment(horizontal='center', vertical='center')
                if i % 2 == 0:
                    cell.fill = alt_fill

        for col in ws.columns:
            max_len = max((len(str(c.value)) if c.value else 0) for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 3, 30)

        ws.row_dimensions[1].height = 30
        ws.freeze_panes = 'B2'

    buf.seek(0)
    return buf

# Session state
if 'alliance_df' not in st.session_state:
    st.session_state.alliance_df = pd.DataFrame()
if 'last_data' not in st.session_state:
    st.session_state.last_data = None

# Step 1: Load existing tracker
st.subheader("1️⃣ Load Existing Tracker (optional)")
existing_file = st.file_uploader("Upload existing Alliance Excel to keep adding players", type=["xlsx"], key="existing")
if existing_file:
    st.session_state.alliance_df = pd.read_excel(existing_file)
    st.success(f"✅ Loaded {len(st.session_state.alliance_df)} existing players.")

# Step 2: Upload screenshots
st.subheader("2️⃣ Upload All Screenshots for One Player")
uploaded_files = st.file_uploader(
    "Select all screenshots at once",
    type=["png", "jpg", "jpeg", "webp"],
    accept_multiple_files=True,
    key="shots"
)

if uploaded_files:
    st.write(f"📸 {len(uploaded_files)} screenshots ready")

    if st.button("🔍 Extract Stats"):
        all_text = ""
        progress = st.progress(0)
        for i, f in enumerate(uploaded_files):
            img = Image.open(f)
            all_text += "\n" + ocr_image(img)
            progress.progress((i + 1) / len(uploaded_files))

        extracted = extract_all(all_text)
        st.session_state.last_data = extracted

        st.subheader("✅ Extracted Data")

        player_keys = ['Player Name', 'Result', 'Power Loss', 'Battlers',
                       'Total Army', 'Kills', 'Losses', 'Wounded', 'Survivors']
        other_keys = ['Elder Titan Tier', 'Titan Talent Level', 'Beast Tier',
                      'Beast Talent Level', 'Beast Skill Level', 'Totem Level',
                      'Special Stats Level', 'Jewels Level', 'Zodiac Green', 'Zodiac White',
                      'Northern Green', 'Colossus Level', 'Emblem Level']

        col1, col2 = st.columns(2)
        with col1:
            st.markdown("**⚔️ Battle Info**")
            battle_data = {k: extracted.get(k, '—') for k in player_keys}
            st.dataframe(pd.DataFrame(list(battle_data.items()), columns=['Field', 'Value']), use_container_width=True)
        with col2:
            st.markdown("**🏆 Player Build**")
            other_data = {k: extracted.get(k, '—') for k in other_keys if k in extracted}
            st.dataframe(pd.DataFrame(list(other_data.items()), columns=['Field', 'Value']), use_container_width=True)

        st.markdown("**📊 Stats Bonus**")
        stats_data = {k: extracted[k] for k in KNOWN_STATS if k in extracted}
        if stats_data:
            items = list(stats_data.items())
            half = len(items) // 2 + len(items) % 2
            c1, c2 = st.columns(2)
            with c1:
                st.dataframe(pd.DataFrame(items[:half], columns=['Stat', 'Value']), use_container_width=True)
            with c2:
                st.dataframe(pd.DataFrame(items[half:], columns=['Stat', 'Value']), use_container_width=True)
            st.success(f"✅ {len(stats_data)}/47 stats extracted successfully!")
        else:
            st.warning("No stats bonus found — check raw text below.")

        with st.expander("🔤 Raw OCR Text (for debugging)"):
            st.text_area("", all_text, height=300)

# Step 3: Add to tracker
if st.session_state.last_data:
    st.subheader("3️⃣ Add to Alliance Tracker")
    player_name = st.session_state.last_data.get('Player Name', 'Unknown Player')
    total_fields = len(st.session_state.last_data)
    st.write(f"Ready to add: **{player_name}** — {total_fields} fields extracted")

    if st.button(f"➕ Add {player_name} to Tracker"):
        new_row = pd.DataFrame([st.session_state.last_data])
        st.session_state.alliance_df = pd.concat(
            [st.session_state.alliance_df, new_row], ignore_index=True
        )
        st.session_state.last_data = None
        total = len(st.session_state.alliance_df)
        st.success(f"✅ Added! Total players in tracker: {total}")

# Step 4: View and download
if not st.session_state.alliance_df.empty:
    st.subheader(f"📊 Alliance Tracker — {len(st.session_state.alliance_df)} Players")
    st.dataframe(st.session_state.alliance_df, use_container_width=True)
    excel_buf = build_excel(st.session_state.alliance_df)
    st.download_button(
        label="⬇️ Download Alliance Tracker Excel",
        data=excel_buf,
        file_name="alliance_tracker.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
